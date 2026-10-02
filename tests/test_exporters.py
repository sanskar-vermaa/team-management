import csv
import io

from openpyxl import load_workbook

from team_builder.exporters import EXPORTERS, to_csv, to_pdf, to_xlsx
from team_builder.teams import TeamPlan

PLAN = TeamPlan(
    teams=[{"Dev": ["Asha", "Ben"], "BA": ["Ravi"]}, {"Dev": ["Chen", "=HYPERLINK(1)"], "BA": ["Meera"]}],
    leftovers={"Dev": ["Zoe"]},
)


def test_csv_lists_every_member_with_team_and_role():
    text = to_csv(PLAN).getvalue().decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == ["Team", "Role", "Member"]
    assert ["Team 1", "Dev", "Asha"] in rows
    assert ["Team 2", "BA", "Meera"] in rows
    assert ["Unassigned", "Dev", "Zoe"] in rows
    assert len(rows) == 1 + 7


def test_csv_neutralises_formulas():
    text = to_csv(PLAN).getvalue().decode("utf-8-sig")
    assert ["Team 2", "Dev", "'=HYPERLINK(1)"] in list(csv.reader(io.StringIO(text)))


def test_csv_starts_with_a_bom_for_excel():
    assert to_csv(PLAN).getvalue().startswith("﻿".encode())


def test_xlsx_has_a_summary_and_one_sheet_per_team():
    wb = load_workbook(to_xlsx(PLAN))
    assert wb.sheetnames == ["All teams", "Team 1", "Team 2"]
    summary = [tuple(r) for r in wb["All teams"].iter_rows(values_only=True)]
    assert summary[0] == ("Team", "Role", "Member")
    assert ("Unassigned", "Dev", "Zoe") in summary
    team2 = [tuple(r) for r in wb["Team 2"].iter_rows(values_only=True)]
    assert ("Dev", "'=HYPERLINK(1)") in team2
    assert wb["All teams"]["A1"].font.bold


def test_pdf_is_a_valid_document():
    data = to_pdf(PLAN).getvalue()
    assert data.startswith(b"%PDF")
    assert data.rstrip().endswith(b"%%EOF")


def test_pdf_escapes_markup_in_the_title():
    assert to_pdf(PLAN, title="<b>Q3 & Q4</b>").getvalue().startswith(b"%PDF")


def test_registry_maps_formats_to_mimetypes():
    assert set(EXPORTERS) == {"pdf", "csv", "xlsx"}
    assert EXPORTERS["pdf"][1] == "application/pdf"
