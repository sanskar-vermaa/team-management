import csv
import io
import json
import re

from openpyxl import load_workbook

from tests.conftest import ROSTER


def flashes(response) -> list[str]:
    return re.findall(r'class="flash [a-z]+" role="alert">([^<]+)<', response.get_data(as_text=True))


def test_home_page(client):
    res = client.get("/")
    assert res.status_code == 200
    assert b"Build teams from a roster" in res.data
    assert b'accept=".xlsx, .xls, .csv"' in res.data


def test_upload_shows_detected_roles_with_defaults(client, roster_xlsx):
    res = client.post("/upload", data={"file": (roster_xlsx(), "team.xlsx")})
    html = res.get_data(as_text=True)
    assert res.status_code == 200
    assert "3 roles" in html and "12 people" in html
    assert 'name="count_0"\n                   value="3"' in html  # Dev defaults to 3
    assert 'name="count_1"\n                   value="1"' in html


def test_upload_csv(client):
    data = io.BytesIO(b"Frontend,Backend\nA,B\nC,D\n")
    res = client.post("/upload", data={"file": (data, "r.csv")})
    assert res.status_code == 200
    assert "Frontend" in res.get_data(as_text=True)


def test_upload_without_a_file(client):
    res = client.post("/upload", data={}, follow_redirects=True)
    assert flashes(res) == ["Please choose a roster file to upload."]


def test_upload_rejects_other_file_types(client):
    res = client.post("/upload", data={"file": (io.BytesIO(b"x"), "notes.pdf")}, follow_redirects=True)
    assert "Unsupported file type" in flashes(res)[0]


def test_upload_reports_duplicates_as_warnings(client, roster_xlsx):
    res = client.post("/upload", data={"file": (roster_xlsx({"Dev": ["Asha", "asha"]}), "t.xlsx")})
    assert res.status_code == 200
    assert flashes(res) == ["Removed duplicate &#39;asha&#39; from Dev."]


def test_generate_teams(client, roster_json):
    res = client.post(
        "/teams", data={"roster_json": roster_json, "count_0": "3", "count_1": "1", "count_2": "1"}
    )
    html = res.get_data(as_text=True)
    assert res.status_code == 200
    assert "2 teams generated" in html
    assert "2 unassigned" in html
    assert "Dev: d7" in html


def test_generate_with_shuffle_shows_a_reproducible_seed(client, roster_json):
    form = {"roster_json": roster_json, "count_0": "3", "count_1": "1", "count_2": "1", "shuffle": "on"}
    first = client.post("/teams", data=form).get_data(as_text=True)
    seed = re.search(r"seed (\d+)", first).group(1)
    again = client.post("/teams", data={**form, "seed": seed}).get_data(as_text=True)
    teams = lambda html: re.findall(r'<div class="member"><span>([^<]+)</span>', html)  # noqa: E731
    assert teams(first) == teams(again)


def test_generate_with_distribution_leaves_nobody_out(client, roster_json):
    res = client.post(
        "/teams",
        data={"roster_json": roster_json, "count_0": "3", "count_1": "1", "count_2": "1", "distribute": "on"},
    )
    html = res.get_data(as_text=True)
    assert "unassigned" not in html
    assert len(re.findall(r'class="member"', html)) == 12


def test_generate_explains_impossible_compositions(client, roster_json):
    res = client.post(
        "/teams", data={"roster_json": roster_json, "count_0": "3", "count_1": "5", "count_2": "1"}
    )
    assert res.status_code == 422
    assert "BA (need 5, have 2)" in flashes(res)[0]
    assert 'value="5"' in res.get_data(as_text=True)  # keeps what the user typed


def test_generate_rejects_non_numeric_counts(client, roster_json):
    res = client.post("/teams", data={"roster_json": roster_json, "count_0": "three"}, follow_redirects=True)
    assert flashes(res) == ["Count for Dev must be a whole number."]


def test_generate_with_tampered_roster(client):
    res = client.post("/teams", data={"roster_json": "{not json"}, follow_redirects=True)
    assert "upload it again" in flashes(res)[0]


def plan_json():
    return json.dumps({"teams": [{"Dev": ["d1", "d2"], "BA": ["b1"]}], "leftovers": {"DA": ["a1"]}})


def test_export_pdf(client):
    res = client.post("/export/pdf", data={"plan_json": plan_json()})
    assert res.status_code == 200
    assert res.mimetype == "application/pdf"
    assert "teams.pdf" in res.headers["Content-Disposition"]
    assert res.data.startswith(b"%PDF")


def test_export_csv(client):
    res = client.post("/export/csv", data={"plan_json": plan_json()})
    rows = list(csv.reader(io.StringIO(res.data.decode("utf-8-sig"))))
    assert rows == [
        ["Team", "Role", "Member"],
        ["Team 1", "Dev", "d1"],
        ["Team 1", "Dev", "d2"],
        ["Team 1", "BA", "b1"],
        ["Unassigned", "DA", "a1"],
    ]


def test_export_xlsx(client):
    res = client.post("/export/xlsx", data={"plan_json": plan_json()})
    assert load_workbook(io.BytesIO(res.data)).sheetnames == ["All teams", "Team 1"]


def test_export_unknown_format_goes_home(client):
    res = client.post("/export/docx", data={"plan_json": plan_json()})
    assert res.status_code == 302


def test_export_rejects_tampered_plans(client):
    res = client.post(
        "/export/pdf", data={"plan_json": json.dumps({"teams": [{"Dev": [1]}]})}, follow_redirects=True
    )
    assert flashes(res) == ["Nothing to export. Generate teams first."]


def test_sample_roster_round_trips(client):
    res = client.get("/sample-roster.xlsx")
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.data))
    assert [c.value for c in wb.active[1]] == ["Dev", "BA", "DA"]

    upload = client.post("/upload", data={"file": (io.BytesIO(res.data), "sample-roster.xlsx")})
    assert "12 people" in upload.get_data(as_text=True)


def test_results_page_offers_every_export(client, roster_json):
    html = client.post(
        "/teams", data={"roster_json": roster_json, "count_0": "3", "count_1": "1", "count_2": "1"}
    ).get_data(as_text=True)
    for fmt in ("pdf", "xlsx", "csv"):
        assert f'action="/export/{fmt}"' in html


def test_full_flow_from_upload_to_export(client, roster_xlsx):
    configure = client.post("/upload", data={"file": (roster_xlsx(), "team.xlsx")}).get_data(as_text=True)
    roster = re.search(r'name="roster_json" value="([^"]+)"', configure).group(1)
    roster = roster.replace("&#34;", '"').replace("&quot;", '"')
    assert json.loads(roster) == ROSTER

    results = client.post(
        "/teams", data={"roster_json": roster, "count_0": "2", "count_1": "1", "count_2": "1"}
    )
    plan = re.search(r'name="plan_json" value="([^"]+)"', results.get_data(as_text=True)).group(1)
    plan = plan.replace("&#34;", '"').replace("&quot;", '"')
    pdf = client.post("/export/pdf", data={"plan_json": plan})
    assert pdf.data.startswith(b"%PDF")
