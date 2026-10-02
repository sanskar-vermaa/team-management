"""Export a team plan as PDF, CSV or Excel."""

from __future__ import annotations

import csv
import io
from datetime import date
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from team_builder.teams import TeamPlan

HEADER = ["Team", "Role", "Member"]


def _rows(plan: TeamPlan):
    for number, team in enumerate(plan.teams, start=1):
        for role, members in team.items():
            for member in members:
                yield [f"Team {number}", role, member]
    for role, members in plan.leftovers.items():
        for member in members:
            yield ["Unassigned", role, member]


def _safe_cell(value: str) -> str:
    # Stop spreadsheet apps from evaluating names like "=HYPERLINK(...)".
    return "'" + value if value[:1] in ("=", "+", "-", "@", "\t", "\r") else value


def to_csv(plan: TeamPlan) -> io.BytesIO:
    text = io.StringIO()
    writer = csv.writer(text)
    writer.writerow(HEADER)
    for row in _rows(plan):
        writer.writerow([_safe_cell(v) for v in row])
    # BOM so Excel opens UTF-8 names (e.g. accents, Devanagari) correctly.
    return io.BytesIO(("﻿" + text.getvalue()).encode("utf-8"))


def to_xlsx(plan: TeamPlan) -> io.BytesIO:
    wb = Workbook()
    summary = wb.active
    summary.title = "All teams"
    summary.append(HEADER)
    for row in _rows(plan):
        summary.append([_safe_cell(v) for v in row])

    for number, team in enumerate(plan.teams, start=1):
        sheet = wb.create_sheet(f"Team {number}")
        sheet.append(["Role", "Member"])
        for role, members in team.items():
            for member in members:
                sheet.append([role, _safe_cell(member)])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F2937")
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(12, width + 2), 50)
        sheet.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def to_pdf(plan: TeamPlan, title: str = "Team Compositions") -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=title)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(escape(title), styles["Title"]),
        Paragraph(
            f"{len(plan.teams)} teams · generated {date.today():%d %b %Y}",
            styles["Normal"],
        ),
        Spacer(1, 16),
    ]

    table_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f4f5")]),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
    )

    def table_for(group: dict[str, list[str]]) -> Table:
        rows = [["Role", "Member"]]
        rows += [[role, member] for role, members in group.items() for member in members]
        table = Table(rows, colWidths=[120, 330], repeatRows=1)
        table.setStyle(table_style)
        return table

    for number, team in enumerate(plan.teams, start=1):
        size = sum(len(m) for m in team.values())
        elements.append(Paragraph(f"Team {number} ({size} people)", styles["Heading2"]))
        elements.append(table_for(team))
        elements.append(Spacer(1, 14))

    if plan.leftovers:
        elements.append(Paragraph("Unassigned", styles["Heading2"]))
        elements.append(table_for(plan.leftovers))

    doc.build(elements)
    buffer.seek(0)
    return buffer


EXPORTERS = {
    "pdf": (to_pdf, "application/pdf"),
    "csv": (to_csv, "text/csv"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}
